"""Safe text extraction from uploaded PDF/DOCX files.

Validation is by content (magic bytes), not just file name. Raises DocumentError with a
user-safe message. Scanned/image-only PDFs raise ScannedDocumentError instead of
returning an empty analysis.
"""
from __future__ import annotations

import io
import zipfile

MAX_PAGES = 12
MAX_DOCX_UNCOMPRESSED = 25 * 1024 * 1024
MIN_TEXT_CHARS = 120


class DocumentError(Exception):
    """Invalid or unreadable document (message is safe to show to the user)."""


class ScannedDocumentError(DocumentError):
    """Document opened fine but contains (almost) no extractable text."""


def detect_type(data: bytes) -> str:
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
    raise DocumentError("Only valid PDF or DOCX files are accepted.")


def _pdf_text(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:  # pragma: no cover
        raise DocumentError("PDF support is not installed on the server.") from e
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as e:
        raise DocumentError("This PDF could not be read. It may be damaged or password-protected.") from e
    with doc:
        if doc.needs_pass:
            raise DocumentError("Password-protected PDFs are not supported.")
        if doc.page_count > MAX_PAGES:
            raise DocumentError(f"Resumes longer than {MAX_PAGES} pages are not supported.")
        return "\n".join(page.get_text("text") for page in doc)


def _docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise DocumentError("This DOCX file is too large once unpacked.")
        import docx
        d = docx.Document(io.BytesIO(data))
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError("This DOCX could not be read. It may be damaged.") from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_text(data: bytes, max_bytes: int) -> tuple[str, str]:
    """Return (text, file_type)."""
    if not data:
        raise DocumentError("The file is empty.")
    if len(data) > max_bytes:
        raise DocumentError(f"The file is larger than {max_bytes // (1024 * 1024)} MB.")
    kind = detect_type(data)
    text = _pdf_text(data) if kind == "pdf" else _docx_text(data)
    text = text.replace("\x00", "").strip()
    if len(text) < MIN_TEXT_CHARS:
        if kind == "pdf":
            raise ScannedDocumentError(
                "No readable text was found. This looks like a scanned or image-only PDF. "
                "Export a text-based PDF from your editor, upload a DOCX, or run OCR first.")
        raise DocumentError("Very little text was found in this document.")
    return text, kind
